"""
tests/test_data_hub_sync_clear.py
─────────────────────────────────
Verifies that triggering Cloud-to-Local synchronization clears existing local table data
from:
1. Exchange Master (kalai_exchangemaster / ExchangeMasterData)
2. Strategy State (kalai_algoinfo / AlgoInfo)
3. Algorithm Logs (kalai_algolog / AlgoLog)
4. Positions & P&L Snapshots (kalai_brokerposition, kalai_dailypnlsnapshot, kalai_traderecord)
"""

import json
from decimal import Decimal
from unittest.mock import patch, MagicMock
import pytest
from django.contrib.auth.models import User
from django.test import Client
from django.utils import timezone
from kalai.models import (
    Broker,
    BrokerType,
    ApiProvider,
    AlgoInfo,
    ExchangeMasterData,
    AlgoLog,
    BrokerPosition,
    DailyPnLSnapshot,
    TradeRecord,
)
from kalai.views import _clear_table_data


@pytest.fixture
def test_broker(db):
    bt, _ = BrokerType.objects.get_or_create(name="INDIAN")
    ap, _ = ApiProvider.objects.get_or_create(name="ZERODHA")
    broker, _ = Broker.objects.get_or_create(
        account_id="TEST_SYNC_ACC",
        defaults={
            "name": "Test Sync Account",
            "broker_name": bt,
            "api_provider": ap,
            "enable_trade": True,
        }
    )
    return broker


@pytest.fixture
def staff_client(db):
    user = User.objects.create_superuser(
        username="sync_admin",
        email="sync_admin@test.com",
        password="testpassword123"
    )
    client = Client()
    client.force_login(user)
    return client


@pytest.mark.django_db
def test_clear_table_data_helper(test_broker):
    """Verify that _clear_table_data empties target tables completely."""
    # 1. Exchange Master
    ExchangeMasterData.objects.create(
        account=test_broker,
        master_name="cum_table",
        row_count=10,
        tabledata=[{"symbol": "NIFTY"}],
    )
    assert ExchangeMasterData.objects.count() == 1
    _clear_table_data(ExchangeMasterData)
    assert ExchangeMasterData.objects.count() == 0

    # 2. Strategy State (AlgoInfo)
    AlgoInfo.objects.create(
        account=test_broker,
        tablename="state_test",
        tabledata={"active": True},
    )
    assert AlgoInfo.objects.count() == 1
    _clear_table_data(AlgoInfo)
    assert AlgoInfo.objects.count() == 0

    # 3. Algorithm Logs
    AlgoLog.objects.create(
        account=test_broker,
        algo_name="test_algo",
        tag="TEST",
        level="INFO",
        message="Test log message",
    )
    assert AlgoLog.objects.count() == 1
    _clear_table_data(AlgoLog)
    assert AlgoLog.objects.count() == 0

    # 4. Positions & P&L Snapshots
    BrokerPosition.objects.create(
        account=test_broker,
        tradingsymbol="NIFTY26SEP24000CE",
        product="NRML",
        quantity=Decimal("50"),
    )
    DailyPnLSnapshot.objects.create(
        account=test_broker,
        date=timezone.now().date(),
        realized_pnl=Decimal("1500.00"),
    )
    TradeRecord.objects.create(
        account=test_broker,
        tradingsymbol="NIFTY26SEP24000CE",
        action_type="BUY",
        quantity=Decimal("50"),
        price=Decimal("120.50"),
    )
    assert BrokerPosition.objects.count() == 1
    assert DailyPnLSnapshot.objects.count() == 1
    assert TradeRecord.objects.count() == 1

    _clear_table_data(BrokerPosition)
    _clear_table_data(DailyPnLSnapshot)
    _clear_table_data(TradeRecord)

    assert BrokerPosition.objects.count() == 0
    assert DailyPnLSnapshot.objects.count() == 0
    assert TradeRecord.objects.count() == 0


@pytest.mark.django_db
def test_export_pnl_api(staff_client, test_broker):
    """Verify /api/export/pnl/ exports position and PnL models in JSON format."""
    BrokerPosition.objects.create(
        account=test_broker,
        tradingsymbol="BANKNIFTY26SEP50000CE",
        product="NRML",
        quantity=Decimal("15"),
    )
    res = staff_client.get("/api/export/pnl/")
    assert res.status_code == 200
    assert res["Content-Type"] == "application/json"
    data = json.loads(res.content)
    assert any(item["model"] == "kalai.brokerposition" for item in data)


@pytest.mark.django_db
@patch("requests.get")
def test_clone_remote_data_clears_tables_on_sync(mock_get, staff_client, test_broker):
    """
    Verify that clone_remote_data_api clears existing local records
    for Exchange Master, Strategy State, AlgoLog, and Positions & P&L before import.
    """
    # Pre-populate records
    ExchangeMasterData.objects.create(account=test_broker, master_name="old_master", row_count=5)
    AlgoInfo.objects.create(account=test_broker, tablename="old_state")
    AlgoLog.objects.create(account=test_broker, algo_name="old_algo", message="old log")
    BrokerPosition.objects.create(account=test_broker, tradingsymbol="OLD_POS", product="NRML", quantity=Decimal("10"))
    DailyPnLSnapshot.objects.create(account=test_broker, date=timezone.now().date(), realized_pnl=Decimal("100"))
    TradeRecord.objects.create(account=test_broker, tradingsymbol="OLD_TRADE", action_type="BUY")

    assert ExchangeMasterData.objects.count() == 1
    assert AlgoInfo.objects.count() == 1
    assert AlgoLog.objects.count() == 1
    assert BrokerPosition.objects.count() == 1
    assert DailyPnLSnapshot.objects.count() == 1
    assert TradeRecord.objects.count() == 1

    # Mock remote responses to return empty or 200 with empty JSON list
    mock_response = MagicMock()
    mock_response.status_code = 200
    mock_response.text = "[]"
    mock_response.content = b""
    mock_get.return_value = mock_response

    payload = {
        "remote_host": "34.180.29.46",
        "hours": 24,
        "chunk_hours": 4,
        "datasets": ["masters", "state", "logs", "pnl"],
        "clear_tables": True,
        "pause": 0.0,
    }

    res = staff_client.post(
        "/api/data-hub/clone/",
        data=json.dumps(payload),
        content_type="application/json",
    )

    assert res.status_code == 200
    res_data = res.json()
    assert res_data["success"] is True

    # Check details indicate tables were cleared
    details = res_data.get("details", {})
    assert "Exchange Masters" in details and "Cleared local table" in details["Exchange Masters"]
    assert "Strategy State (AlgoInfo)" in details and "Cleared local table" in details["Strategy State (AlgoInfo)"]
    assert "Algorithm Logs" in details and "Cleared local table" in details["Algorithm Logs"]
    assert "Positions & P&L Snapshots" in details and "Cleared local tables" in details["Positions & P&L Snapshots"]

    # Verify existing old records are gone
    assert ExchangeMasterData.objects.count() == 0
    assert AlgoInfo.objects.count() == 0
    assert AlgoLog.objects.count() == 0
    assert BrokerPosition.objects.count() == 0
    assert DailyPnLSnapshot.objects.count() == 0
    assert TradeRecord.objects.count() == 0


@pytest.mark.django_db
@patch("kalai.views.requests.get")
def test_clone_remote_data_respects_ssl_verify_flag(mock_get, staff_client, test_broker):
    """Verify that ssl_verify=False in payload passes verify=False to requests.get."""
    mock_response = MagicMock()
    mock_response.status_code = 200
    mock_response.text = "[]"
    mock_response.content = b""
    mock_get.return_value = mock_response

    payload = {
        "remote_host": "172.235.29.89",
        "hours": 1,
        "chunk_hours": 1,
        "datasets": ["masters"],
        "ssl_verify": False,
        "pause": 0.0,
    }

    res = staff_client.post(
        "/api/data-hub/clone/",
        data=json.dumps(payload),
        content_type="application/json",
    )

    assert res.status_code == 200
    assert mock_get.called
    # Check that verify was False
    _, kwargs = mock_get.call_args
    assert kwargs.get("verify") is False


@pytest.mark.django_db
@patch("kalai.views.requests.get")
def test_clone_remote_data_ssl_error_handling(mock_get, staff_client, test_broker):
    """Verify that SSLError returns a clear 400 response with helpful guidance."""
    import requests
    mock_get.side_effect = requests.exceptions.SSLError("certificate verify failed: self-signed certificate")

    payload = {
        "remote_host": "172.235.29.89",
        "hours": 1,
        "chunk_hours": 1,
        "datasets": ["masters"],
        "ssl_verify": True,
        "pause": 0.0,
    }

    res = staff_client.post(
        "/api/data-hub/clone/",
        data=json.dumps(payload),
        content_type="application/json",
    )

    assert res.status_code == 400
    data = res.json()
    assert data["success"] is False
    assert "SSL Certificate Verification Failed" in data["error"]
    assert "Accept Self-Signed SSL" in data["error"]


@pytest.mark.django_db
@patch("kalai.views.requests.get")
def test_clone_remote_data_foreign_key_remapping(mock_get, staff_client, test_broker):
    """
    Verify that remote records referencing non-existent foreign key account IDs (e.g. account=999)
    are safely remapped to local brokers without violating foreign key constraints.
    """
    def mock_requests_get(url, **kwargs):
        resp = MagicMock()
        resp.status_code = 200
        if "api/export/models" in url:
            resp.content = json.dumps([
                {
                    "model": "kalai.algoinfo",
                    "pk": 1,
                    "fields": {
                        "account": 999,  # Remote broker ID does not exist locally
                        "tablename": f"{test_broker.account_id}_inst_tokens",
                        "tabledata": {"tokenid": [12345]},
                    }
                }
            ]).encode("utf-8")
        elif "api/export/masters" in url:
            resp.content = json.dumps([
                {
                    "model": "kalai.exchangemasterdata",
                    "pk": 1,
                    "fields": {
                        "account": 999,  # Remote broker ID does not exist locally
                        "master_name": "cum_table",
                        "row_count": 1,
                        "tabledata": [{"Symbol": "NIFTY", "strike": 24000}],
                    }
                }
            ]).encode("utf-8")
        else:
            resp.content = b"[]"
        return resp

    mock_get.side_effect = mock_requests_get

    payload = {
        "remote_host": "172.235.29.89",
        "hours": 1,
        "chunk_hours": 1,
        "datasets": ["masters", "state"],
        "ssl_verify": False,
        "clear_tables": True,
        "pause": 0.0,
    }

    res = staff_client.post(
        "/api/data-hub/clone/",
        data=json.dumps(payload),
        content_type="application/json",
    )

    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True

    # Verify that records were saved and assigned to test_broker (not account_id=999)
    saved_master = ExchangeMasterData.objects.filter(master_name="cum_table").first()
    assert saved_master is not None
    assert saved_master.account == test_broker

    saved_state = AlgoInfo.objects.filter(tablename=f"{test_broker.account_id}_inst_tokens").first()
    assert saved_state is not None
    assert saved_state.account == test_broker


def test_export_cum_table_excel_exact_tabledata(db, staff_client, test_broker):
    """Verify that exporting exchangemaster in Excel produces actual contract tabledata, not just file metadata."""
    import openpyxl

    ExchangeMasterData.objects.create(
        account=test_broker,
        master_name="cum_table",
        row_count=2,
        tabledata=[
            {"tradingsymbol": "NIFTY24SEP26000CE", "strike": 26000.0, "lot_size": 75, "instrument_type": "CE"},
            {"tradingsymbol": "NIFTY24SEP26000PE", "strike": 26000.0, "lot_size": 75, "instrument_type": "PE"},
        ],
    )

    res = staff_client.get(
        f"/api/data-hub/export/?table=exchangemaster&account_id={test_broker.account_id}&format=xlsx",
        HTTP_HOST="localhost",
    )
    assert res.status_code == 200
    assert "cum_table" in res.headers.get("Content-Disposition", "")

    # Load the resulting workbook
    import io
    wb = openpyxl.load_workbook(io.BytesIO(b"".join(res.streaming_content)))
    assert "CUM_TABLE" in wb.sheetnames
    assert wb.active.title == "CUM_TABLE"
    ws = wb["CUM_TABLE"]

    # Verify header row + 2 contract rows
    assert ws.max_row == 3
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert "Account_ID" in headers
    assert "tradingsymbol" in headers
    assert "strike" in headers

    # Verify actual row values
    row1_vals = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    assert test_broker.account_id in row1_vals
    assert "NIFTY24SEP26000CE" in row1_vals
    assert 26000.0 in row1_vals


def test_export_cum_table_csv_exact_tabledata(db, staff_client, test_broker):
    """Verify that exporting exchangemaster in CSV produces actual contract tabledata rows."""
    ExchangeMasterData.objects.create(
        account=test_broker,
        master_name="cum_table",
        row_count=1,
        tabledata=[
            {"tradingsymbol": "BTC24SEP60000CE", "strike": 60000.0, "lot_size": 0.001, "instrument_type": "CE"},
        ],
    )

    res = staff_client.get(
        f"/api/data-hub/export/?table=exchangemaster&account_id={test_broker.account_id}&format=csv",
        HTTP_HOST="localhost",
    )
    assert res.status_code == 200
    content = res.content.decode("utf-8")
    assert "tradingsymbol" in content
    assert "BTC24SEP60000CE" in content
    assert "60000.0" in content



